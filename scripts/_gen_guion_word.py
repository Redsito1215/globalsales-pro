# -*- coding: utf-8 -*-
"""Genera el Word de guiones operativo / táctico / estratégico."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "Guion-sistemas-operativo-tactico-estrategico.docx"

PURPLE = RGBColor(0x71, 0x4B, 0x67)
DARK = RGBColor(0x1F, 0x1A, 0x24)
MUTED = RGBColor(0x5A, 0x54, 0x62)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, *, size=11, bold=False, italic=False, color=DARK, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, italic=False, color=DARK, space_after=8, space_before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_mixed(doc, parts, *, space_after=8, space_before=0):
    """parts: list of (text, kwargs)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run(run, **kwargs)
    return p


def shade_heading(paragraph, fill="714B67"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): fill,
    })
    pPr.append(shd)
    paragraph.paragraph_format.left_indent = Cm(-0.5)
    paragraph.paragraph_format.right_indent = Cm(-0.5)
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.paragraph_format.space_after = Pt(10)
    for run in paragraph.runs:
        run.font.color.rgb = WHITE


def add_block_heading(doc, text):
    p = add_para(doc, "  " + text + "  ", size=16, bold=True, color=WHITE, space_after=10, space_before=16)
    shade_heading(p)
    return p


def add_sub(doc, text):
    return add_para(doc, text, size=13, bold=True, color=PURPLE, space_before=12, space_after=6)


def add_quote(doc, label, text):
    add_mixed(
        doc,
        [
            (label + " ", {"size": 11, "bold": True, "color": PURPLE}),
            ("«" + text + "»", {"size": 11, "italic": True, "color": DARK}),
        ],
        space_after=8,
    )


def add_do(doc, text):
    add_mixed(
        doc,
        [
            ("Haz: ", {"size": 11, "bold": True, "color": PURPLE}),
            (text, {"size": 11, "color": DARK}),
        ],
        space_after=4,
    )


def add_say(doc, text):
    add_mixed(
        doc,
        [
            ("Di: ", {"size": 11, "bold": True, "color": PURPLE}),
            ("«" + text + "»", {"size": 11, "italic": True, "color": DARK}),
        ],
        space_after=10,
    )


def set_cell(cell, text, *, bold=False, fill=None, color=DARK, size=10, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=WHITE if fill in ("714B67", "5C3D56") else color)
    if fill:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): fill,
        })
        tcPr.append(shd)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("GLOBTRADE  ·  Guion de defensa  ·  50 minutos")
    set_run(hr, size=9, italic=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Sistemas de información  ·  Operativo 20 min  ·  Táctico 15 min  ·  Estratégico 15 min")
    set_run(fr, size=8, color=MUTED)

    add_para(doc, "GLOBTRADE", size=12, bold=True, color=PURPLE, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        "Guion de demostración",
        size=22,
        bold=True,
        color=DARK,
        space_after=4,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Sistema operativo (20 min)  ·  Sistema táctico (15 min)  ·  Sistema estratégico (15 min)",
        size=12,
        italic=True,
        color=MUTED,
        space_after=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_para(
        doc,
        "Tres guiones para leer en voz alta. Entra como administrador antes de empezar. "
        "No sigas el menú de carpetas: el Tablero es estratégico aunque esté en Operacionales; "
        "Informes simples son operativos aunque estén en Estratégicas.",
        size=11,
        space_after=12,
    )

    add_sub(doc, "Cómo usar este documento")
    bullets = [
        "Tiempo total: 50 minutos (20 + 15 + 15). Deja 30–40 segundos de colchón al final de cada bloque.",
        "Orden: operativo → táctico → estratégico. No inviertas el orden: el ELT del último bloque necesita el puente que explicaste en el primero.",
        "Si una pantalla falla, usa el Plan B del final. Los tres niveles se sostienen igual.",
        "Frases entre comillas se leen tal cual. Lo marcado como «Haz» son clics; no lo narres palabra por palabra.",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(b)
        set_run(run, size=11)

    # ─── 1 OPERATIVO ───
    add_block_heading(doc, "1.  Sistema operativo  —  20 minutos")
    add_para(
        doc,
        "Nivel: transaccional (TPS / OLTP). Datos en globtrade_ops. Pregunta: qué hay que hacer hoy.",
        size=11,
        italic=True,
        color=MUTED,
        space_after=10,
    )
    add_quote(
        doc,
        "Apertura (30 s)",
        "Este bloque es el sistema transaccional. Aquí se captura y controla el día a día: tienda, pedidos, inventario, compras, soporte. Los datos viven en la base operativa. No hay agregaciones gerenciales: hay transacciones.",
    )

    add_sub(doc, "0:00–2:00  ·  Quién entra y con qué permiso")
    add_do(doc, "Barra superior → iniciar sesión (administrador). Menciona roles: visitante lee, analista exporta, admin escribe.")
    add_say(doc, "Sin sesión, la vitrina se ve. Aprobar, convertir, ajustar stock y cargar datos quedan bloqueados. Eso es control operativo, no estrategia.")

    add_sub(doc, "2:00–6:00  ·  Vitrina y pedido del cliente")
    add_do(doc, "Operacionales → Tienda. Entra a un producto, agrégalo, confirma la solicitud. Luego Mis pedidos.")
    add_say(doc, "El cliente no toca el data warehouse. Genera una solicitud de compra. El estado arranca en pendiente. El ciclo es: solicitar → revisar → pagar o crédito → convertir en venta → enviar → entregar.")

    add_sub(doc, "6:00–11:00  ·  Bandeja comercial")
    add_do(doc, "Pedidos y ventas → pestaña Solicitudes. Filtra «Pendientes de gestión». Abre una: revisar, aprobar, marcar pago o crédito, Convertir. Si hay tiempo, Enviar.")
    add_say(doc, "Esta bandeja es TPS. Cada fila es un documento vivo. Convertir escribe en landing sales_records, todavía no en el tablero. El KPI ejecutivo no se mueve hasta el ELT. Eso separa operación de estrategia.")

    add_sub(doc, "11:00–15:30  ·  Compras e inventario")
    add_do(doc, "Compras e inventario. 1) Pestaña Inventario: bodega general, marca «Solo stock bajo». 2) Proveedores: muestra continente y país. 3) Órdenes de compra: flujo borrador → enviar → recibir.")
    add_say(doc, "Reposición del día: umbral, proveedor, OC, recepción en bodega única. Quien opera stock no espera un informe mensual.")

    add_sub(doc, "15:30–18:00  ·  Soporte y notificaciones")
    add_do(doc, "Soporte (hilo cliente/staff). Notificaciones (avisos de pedido, compra, chat).")
    add_say(doc, "Incidencias y alertas operativas. El correo real está fuera de alcance a propósito: el flujo queda modelado en base y UI.")

    add_sub(doc, "18:00–20:00  ·  Cierre operativo con informe simple")
    add_do(doc, "Estratégicas → Informes simples. Corre RS-01 (solicitudes pendientes) y RS-03 (stock bajo).")
    add_say(doc, "RS son listados del OLTP, no rankings. Cierran el sistema operativo: qué hay que aprobar hoy y qué hay que reponer hoy.")

    add_quote(
        doc,
        "Cierre",
        "Sistema operativo: transacción, estado, documento. Siguiente nivel: el mando medio deja de ver filas y empieza a ver periodos.",
    )

    add_para(doc, "Informes simples de apoyo (si preguntan)", size=11, bold=True, color=PURPLE, space_before=6, space_after=4)
    rs = [
        "RS-01  Listado de solicitudes pendientes de revisión y aprobación",
        "RS-02  Pedidos con pago pendiente o a crédito",
        "RS-03  Productos con stock bajo el mínimo",
        "RS-04  Chats de clientes que esperan respuesta",
        "RS-05  Órdenes de compra abiertas",
        "RS-07  Pedidos listos que aún no se han enviado",
    ]
    for item in rs:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_run(run, size=10)

    # ─── 2 TÁCTICO ───
    add_block_heading(doc, "2.  Sistema táctico  —  15 minutos")
    add_para(
        doc,
        "Nivel: mando medio (MIS). Análisis por periodo. Pregunta: qué región, canal o categoría ajustar este trimestre.",
        size=11,
        italic=True,
        color=MUTED,
        space_after=10,
    )
    add_quote(
        doc,
        "Apertura (20 s)",
        "Sistema táctico: jefe comercial, inventario, logística. Pregunta de semanas y meses: qué región, canal o categoría ajustar. No es el pedido de hoy ni el objetivo a 2030.",
    )

    add_sub(doc, "0:00–3:00  ·  Exploración filtrada")
    add_do(doc, "Estratégicas → Explorar ventas. Filtra país, producto, canal, prioridad. Mira 2 o 3 filas.")
    add_say(doc, "Esto ya no es la bandeja de solicitudes. Es el histórico comercial para comparar. El táctico no crea el pedido: lo lee por criterio.")

    add_sub(doc, "3:00–7:00  ·  Tres cortes de mando medio")
    add_do(doc, "En este orden: 1) Tendencia de ventas — ingresos y utilidad por mes. 2) Ventas por región — dónde concentrar o recortar. 3) Ventas por categoría — mix y rentabilidad.")
    add_say(doc, "Táctico típico: Europa vs Asia este periodo; aperitivos vs cosmética. Decisión: más stock en la categoría que rota, menos esfuerzo en la que no. Horizonte: trimestre, no el turno.")

    add_sub(doc, "7:00–10:00  ·  Catálogo y maestros")
    add_do(doc, "Catálogo analítico (30 s). Luego Gestión: Regiones, Países, Canales.")
    add_say(doc, "Sin maestros limpios el informe táctico miente. Gestión no vende: gobierna dimensiones para que región y canal significen lo mismo en todos los reportes.")

    add_sub(doc, "10:00–13:00  ·  Exportar y un informe de control")
    add_do(doc, "Exportar datos (pide sesión analista o admin). Si hay tiempo, RS-05 (OC abiertas) o RS-08 (enviados sin entregar).")
    add_say(doc, "El CSV sale del sistema hacia Excel o la reunión semanal. RS de logística o compras es táctico corto: cartera de OC y pedidos en tránsito, no el KPI de margen a 7 años.")

    add_sub(doc, "13:00–15:00  ·  Cierre táctico (objetivos)")
    add_do(doc, "Administración → Empresa. Señala Objetivos tácticos: digitalizar pedidos, bajar costo logístico 15 %, CRM/ERP, 50 proveedores, canal online.")
    add_say(doc, "Estos objetivos se miden con tendencias, regiones y mix, no con una solicitud suelta ni con la visión 2030. Puente: si gerencia quiere ver el mismo negocio en una estrella, hay que construir el modelo estratégico.")

    add_quote(
        doc,
        "Cierre",
        "Sistema táctico: periodo, comparación, acción de mando medio. Siguiente: dirección, hechos fact_ventas y decisión de portafolio.",
    )

    # ─── 3 ESTRATÉGICO ───
    add_block_heading(doc, "3.  Sistema estratégico  —  15 minutos")
    add_para(
        doc,
        "Nivel: dirección (EIS / DSS). Modelo estrella fact_ventas + dim_*. Pregunta: si el negocio va hacia las metas de empresa.",
        size=11,
        italic=True,
        color=MUTED,
        space_after=10,
    )
    add_quote(
        doc,
        "Apertura (25 s)",
        "Sistema estratégico: dirección. No lee purchase_requests. Lee el modelo estrella: fact_ventas más dimensiones. Horizonte: histórico 2010–2017 y metas de empresa. DSS/EIS, no TPS.",
    )

    add_sub(doc, "0:00–3:30  ·  El puente ELT (obligatorio)")
    add_do(doc, "Administración → Modelo de datos (ops vs DW, 20 s). Luego Carga ELT. Menciona el DAG globtrade_strategic_etl (Airflow :8080) sin ejecutarlo si el rebuild tarda 15–40 s.")
    add_say(doc, "Pedido convertido queda en landing con analytics desfasado. El tablero no usa operativo. Admin ejecuta Construir modelo: CSV → Parquet → sales_records → fact_ventas. Hasta aquí no hay estrategia; hay bodega.")

    add_sub(doc, "3:30–8:00  ·  Tablero ejecutivo")
    add_do(doc, "Operacionales → Tablero (aclara: menú operativo, capa estratégica). 1) Periodo Todo el histórico. 2) KPIs: pedidos, ingresos, utilidad, costos, margen, países, productos. 3) Un filtro: región Europa o canal En línea. 4) Un gráfico de región o tendencia.")
    add_say(doc, "Vista de gerencia general. Un número, no mil solicitudes. Si el margen baja en una región, la decisión es de portafolio o de presencia, no de aprobar el pedido 128.")

    add_sub(doc, "8:00–12:00  ·  Informes compuestos y Decisiones")
    add_do(doc, "Informes compuestos. Corre RC-01 (ventas mes × categoría) y RC-07 (margen en el tiempo). Si alcanza, RC-02 (top/bottom). Luego Decisiones: margen alto/bajo, stock crítico, embudo, alertas.")
    add_say(doc, "RC agrega la estrella. Decisiones recomienda: qué categoría sostener, qué reponer, dónde se traba el embudo. Eso es DSS: señal para decidir, no listado para operar.")

    add_sub(doc, "12:00–15:00  ·  Cierre con Empresa")
    add_do(doc, "Empresa. Misión, visión, objetivos estratégicos (200+ países, USD 2.000 M, margen ≥ 30 %). Un problema gerencial (por ejemplo sin RFM o sin rentabilidad por SLA).")
    add_say(doc, "El estratégico responde si vamos hacia esas metas. El operativo procesó el pedido. El táctico comparó la región. Este nivel pregunta si el negocio es el correcto.")

    add_quote(
        doc,
        "Cierre",
        "Tres sistemas, una plataforma: OLTP en globtrade_ops, análisis de mando medio sobre el histórico, EIS sobre fact_ventas. El menú mezcla etiquetas; las capas de datos no.",
    )

    add_para(doc, "Informes compuestos de apoyo (si preguntan)", size=11, bold=True, color=PURPLE, space_before=6, space_after=4)
    rc = [
        "RC-01  Ventas por mes y por categoría de producto",
        "RC-02  Top 10 que más se venden y top 10 que menos se venden",
        "RC-07  Ganancia (margen) por categoría en el tiempo",
        "RC-08  Estado de la carga de datos para informes",
    ]
    for item in rc:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        set_run(run, size=10)

    # ─── ORDEN DEL DÍA ───
    add_block_heading(doc, "Orden del día  —  50 minutos")

    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    headers = ["Bloque", "Min", "Entras por", "No hagas aquí"]
    rows = [
        ["Operativo", "20", "Tienda → Ventas → Compras → RS-01 / RS-03", "Tablero, RC, ELT"],
        ["Táctico", "15", "Explorar ventas → Tendencia / Región / Categoría → Exportar", "Convertir pedidos, Construir modelo"],
        ["Estratégico", "15", "ELT → Tablero → RC-01 / RC-07 → Decisiones → Empresa", "Crear OC ni chat de soporte"],
    ]
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, fill="714B67", size=10, center=True)
    fills = ("F7F2F6", "FFFFFF", "F7F2F6")
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell(table.rows[r + 1].cells[c], val, fill=fills[r], size=10, bold=(c == 0))

    add_para(doc, "", space_after=6)

    add_sub(doc, "Plan B si algo falla")
    add_para(
        doc,
        "Operativo = RS-01 + inventario.  Táctico = tendencias + una región.  Estratégico = tablero con «Todo el histórico» + Empresa. "
        "Con eso ya se sostienen los tres niveles ante el jurado.",
        size=11,
        space_after=10,
    )

    add_sub(doc, "Recordatorio si preguntan por el menú")
    add_para(
        doc,
        "Hay dos bases (globtrade_ops y globtrade_dw), no tres. El táctico no es una tercera base: es la capa analítica intermedia "
        "(tendencias, regiones, categorías, exportar, maestros). El Tablero está en el menú Operacionales pero lee fact_ventas. "
        "Los Informes simples están en el menú Estratégicas pero listan el OLTP. En el guion se sigue el nivel gerencial, no la carpeta.",
        size=11,
        space_after=8,
    )

    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    build()
